"""basic_file_reader 单元测试。"""

from __future__ import annotations

import zipfile
import pytest
from pathlib import Path

from src.service.basic_file_reader import (
    BasicFileCategory,
    DASHSCOPE_MAX_MULTIMODAL_BYTES,
    categorize_file,
    estimate_base64_decoded_bytes,
    extract_document_text,
    is_multimodal_payload_too_large,
    read_basic_file,
    read_text_with_encoding_fallback,
)


def test_categorize_file_types() -> None:
    assert categorize_file("a.txt") == BasicFileCategory.TEXT
    assert categorize_file("a.png") == BasicFileCategory.IMAGE
    assert categorize_file("a.pdf") == BasicFileCategory.DOCUMENT
    assert categorize_file("a.docx") == BasicFileCategory.DOCUMENT
    assert categorize_file("a.xlsx") == BasicFileCategory.DOCUMENT


def test_read_text_and_image(tmp_path: Path) -> None:
    text_file = tmp_path / "note.txt"
    text_file.write_text("hello", encoding="utf-8")
    text_payload = read_basic_file(text_file)
    assert text_payload.category == BasicFileCategory.TEXT
    assert text_payload.text == "hello"

    png_bytes = (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01"
        b"\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15\xc4\x89"
        b"\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05\x00\x01\r\n-\xdb"
        b"\x00\x00\x00\x00IEND\xaeB`\x82"
    )
    image_file = tmp_path / "pixel.png"
    image_file.write_bytes(png_bytes)
    image_payload = read_basic_file(image_file)
    assert image_payload.category == BasicFileCategory.IMAGE
    assert image_payload.base64_data
    assert image_payload.mime_type == "image/png"


def test_extract_docx_text(tmp_path: Path) -> None:
    from docx import Document

    docx_path = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("你好，Word")
    document.save(str(docx_path))

    text = extract_document_text(docx_path)
    assert "你好，Word" in text


def test_extract_xlsx_text(tmp_path: Path) -> None:
    from openpyxl import Workbook

    xlsx_path = tmp_path / "sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "数据"
    sheet["A1"] = "名称"
    sheet["B1"] = "数量"
    sheet["A2"] = "苹果"
    sheet["B2"] = 3
    workbook.save(str(xlsx_path))
    workbook.close()

    text = extract_document_text(xlsx_path)
    assert "数据" in text
    assert "苹果" in text


def test_extract_pptx_text(tmp_path: Path) -> None:
    pptx_path = tmp_path / "sample.pptx"
    slide_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
       xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
  <p:cSld><p:spTree><p:sp><p:txBody><a:p><a:r><a:t>演示标题</a:t></a:r></a:p></p:txBody></p:sp></p:spTree></p:cSld>
</p:sld>"""
    content_types = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/ppt/slides/slide1.xml"
    ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>
</Types>"""
    with zipfile.ZipFile(pptx_path, "w") as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("ppt/slides/slide1.xml", slide_xml)

    text = extract_document_text(pptx_path)
    assert "演示标题" in text


def test_multimodal_size_helpers() -> None:
    assert estimate_base64_decoded_bytes("aGVsbG8=") == 5
    assert not is_multimodal_payload_too_large(raw_bytes=1024)
    assert is_multimodal_payload_too_large(
        raw_bytes=DASHSCOPE_MAX_MULTIMODAL_BYTES + 1
    )


def test_read_oversized_image_rejected(tmp_path: Path) -> None:
    huge = tmp_path / "big.png"
    huge.write_bytes(b"\x00" * (DASHSCOPE_MAX_MULTIMODAL_BYTES + 1))
    with pytest.raises(ValueError, match="超过当前模型多模态上限"):
        read_basic_file(huge)


def test_read_text_with_encoding_fallback_gbk(tmp_path: Path) -> None:
    gbk_file = tmp_path / "gbk.md"
    gbk_file.write_bytes("中文测试内容".encode("gbk"))
    assert read_text_with_encoding_fallback(gbk_file) == "中文测试内容"


def test_read_basic_file_gbk_text(tmp_path: Path) -> None:
    gbk_file = tmp_path / "note.md"
    gbk_file.write_bytes("你好 GBK".encode("gbk"))
    payload = read_basic_file(gbk_file)
    assert payload.text == "你好 GBK"
